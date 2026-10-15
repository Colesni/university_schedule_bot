import io
import logging
import re
import csv
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import openpyxl

from .models import ScheduleEntry

logger = logging.getLogger(__name__)

_DATE_RE = re.compile(r"(?<!\d)(\d{1,2})[./-](\d{1,2})[./-](\d{4})(?!\d)")
_TIME_RE = re.compile(r"(?<!\d)(\d{1,2})\s*:\s*(\d{2})\s*[-–—]\s*(\d{1,2})\s*:\s*(\d{2})(?!\d)")
_SUBGROUP_RE = re.compile(r"^\s*(I|II)\s*:\s*(.*)$", re.IGNORECASE)
_TYPE_RE = re.compile(r"\(([^()\n]+)\)\s*,?\s*$")


def _text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"[ \t]+", " ", str(value).replace("\r", "\n")).strip()


def _date_from(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    match = _DATE_RE.search(_text(value))
    if not match:
        return None
    try:
        return date(int(match.group(3)), int(match.group(2)), int(match.group(1)))
    except ValueError:
        return None


def _time_range(value: Any) -> tuple[time, time] | None:
    match = _TIME_RE.search(_text(value))
    if not match:
        return None
    start = time(int(match.group(1)), int(match.group(2)))
    end = time(int(match.group(3)), int(match.group(4)))
    return (start, end) if start < end else None


def _looks_like_group(value: Any) -> bool:
    text = _text(value)
    if not text or "\n" in text or any(word in text.lower() for word in ("anul", "student", "subgrup")):
        return False
    return bool(re.search(r"[A-Za-zĂÂÎȘȚăâîșț]", text) and re.search(r"\d", text))


def _find_group_header(ws: Any) -> tuple[int, dict[int, str]]:
    best: tuple[int, dict[int, str]] | None = None
    for row in range(1, min(ws.max_row, 15) + 1):
        found = {
            cell.column: _text(cell.value)
            for cell in ws[row]
            if _looks_like_group(cell.value)
        }
        if len(found) >= 2 and (best is None or len(found) > len(best[1])):
            best = (row, found)
    if best is None:
        raise ValueError("Could not detect a row containing schedule groups")
    return best


def _find_data_columns(ws: Any, start_row: int, group_columns: dict[int, str]) -> tuple[int, int]:
    date_scores: dict[int, int] = {}
    time_scores: dict[int, int] = {}
    for col in range(1, ws.max_column + 1):
        for row in range(start_row, ws.max_row + 1):
            value = ws.cell(row, col).value
            if _date_from(value):
                date_scores[col] = date_scores.get(col, 0) + 1
            if _time_range(value):
                time_scores[col] = time_scores.get(col, 0) + 1
    date_col = max(date_scores, key=date_scores.get, default=0)
    time_col = max(time_scores, key=time_scores.get, default=0)
    if not date_col or not time_col:
        raise ValueError("Could not detect date and time columns")
    return date_col, time_col


def _split_activity_blocks(value: str) -> list[tuple[str | None, list[str]]]:
    lines = [line.strip() for line in value.replace("\r", "\n").split("\n") if line.strip()]
    blocks: list[tuple[str | None, list[str]]] = []
    subgroup: str | None = None
    current: list[str] = []
    for line in lines:
        marker = _SUBGROUP_RE.match(line)
        if marker:
            if current:
                blocks.append((subgroup, current))
            subgroup = marker.group(1).upper()
            current = [marker.group(2).strip()]
        elif current and len(current) >= 2 and _TYPE_RE.search(line):
            blocks.append((subgroup, current))
            current = [line]
        else:
            current.append(line)
    if current:
        blocks.append((subgroup, current))
    if not blocks and lines:
        blocks.append((None, lines))
    return blocks


def _parse_activity(block: list[str]) -> tuple[str, str | None, str | None, str | None] | None:
    if not block:
        return None
    subject_line = block[0].strip().rstrip(",")
    type_match = _TYPE_RE.search(subject_line)
    activity_type = type_match.group(1).strip() if type_match else None
    if type_match:
        subject_line = subject_line[: type_match.start()].rstrip(" ,")
    if not subject_line:
        return None
    teacher = room = None
    if len(block) > 1:
        details = ", ".join(part.strip() for part in block[1:] if part.strip())
        pieces = [part.strip() for part in details.split(",") if part.strip()]
        if len(pieces) >= 2:
            teacher, room = ", ".join(pieces[:-1]), pieces[-1]
        elif pieces:
            teacher = pieces[0]
    return subject_line, activity_type, teacher, room


class ScheduleParser:
    def parse(self, workbook: Any) -> list[ScheduleEntry]:
        ws = workbook["orar"] if "orar" in workbook.sheetnames else workbook.active
        header_row, group_columns = _find_group_header(ws)
        date_col, time_col = _find_data_columns(ws, header_row + 1, group_columns)
        entries: list[ScheduleEntry] = []
        current_date: date | None = None
        for row in range(header_row + 1, ws.max_row + 1):
            row_date = _date_from(ws.cell(row, date_col).value)
            if row_date:
                current_date = row_date
            time_range = _time_range(ws.cell(row, time_col).value)
            if not current_date or not time_range:
                continue
            for col, group in group_columns.items():
                value = _text(ws.cell(row, col).value)
                if not value:
                    continue
                for subgroup, block in _split_activity_blocks(value):
                    parsed = _parse_activity(block)
                    if not parsed:
                        logger.warning("Skipping invalid activity at row=%s column=%s", row, col)
                        continue
                    subject, activity_type, teacher, room = parsed
                    entries.append(
                        ScheduleEntry(
                            group=group,
                            date=current_date,
                            start_time=time_range[0],
                            end_time=time_range[1],
                            subject=subject,
                            activity_type=activity_type,
                            teacher=teacher,
                            room=room,
                            subgroup=subgroup,
                        )
                    )
        return entries


def parse_workbook(source: str | Path | bytes | bytearray | io.BufferedIOBase) -> list[ScheduleEntry]:
    if isinstance(source, (str, Path)):
        workbook = openpyxl.load_workbook(source, data_only=True)
    else:
        raw = bytes(source) if isinstance(source, (bytes, bytearray)) else source.read()
        if raw[:2] != b"PK":
            rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
            workbook = openpyxl.Workbook()
            ws = workbook.active
            ws.title = "orar"
            for row in rows:
                ws.append(row)
        else:
            workbook = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    return ScheduleParser().parse(workbook)
